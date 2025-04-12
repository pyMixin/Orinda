import os
import json
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter.scrolledtext import ScrolledText
import tkinter.font as tkFont
from PIL import Image, ImageTk  # For handling the logo
from ollama import chat, ChatResponse  # Ollama LLM client

# For file text extraction:
import docx              # For .docx files (python-docx)
import openpyxl          # For .xlsx files
import PyPDF2            # For PDF files

# For Chroma DB integration
import chromadb

# NOTE: Remove 'persist_directory' from the client creation.
# If needed, set the environment variable accordingly.
# os.environ["CHROMA_DB_PERSIST_DIRECTORY"] = "./chroma_data"

# ============================================================
# Global Font Settings (Open Sans, 15pt)
# ============================================================
BASE_FONT = ("Open Sans", 15)

# ============================================================
# LLM Chat Function using Ollama (max_tokens removed)
# ============================================================
def llm_chat(query):
    """
    Process a chat query using the Ollama client with the 'llama3.2' model.
    """
    try:
        response: ChatResponse = chat(
            model="llama3.2",
            messages=[{"role": "user", "content": query}]
        )
        return response["message"]["content"]
    except Exception as e:
        return f"Error calling Ollama: {e}"

# ============================================================
# Persistence Layer (Placeholder for agent config)
# ============================================================
def save_agent_config(agent_config, filename="agent_configs.json"):
    data = {}
    if os.path.exists(filename):
        try:
            with open(filename, "r") as f:
                data = json.load(f)
        except Exception as e:
            print("Error loading existing configuration:", e)
    agent_name = agent_config.get("name", "default_agent")
    data[agent_name] = agent_config
    with open(filename, "w") as f:
        json.dump(data, f, indent=4)
    print(f"Agent configuration for '{agent_name}' saved.")

# ============================================================
# Business Logic: Embedding Functions (Dummy Implementation)
# ============================================================
def generate_embedding(text):
    """
    Dummy function to generate an embedding for the given text.
    Replace with your actual embedding generator.
    Returns a dummy 768-dimensional vector.
    """
    dummy_vector = [0.1] * 768
    return dummy_vector

# ============================================================
# File Extraction Functions
# ============================================================
def extract_text_from_docx(filepath):
    try:
        doc = docx.Document(filepath)
        full_text = [para.text for para in doc.paragraphs]
        return "\n".join(full_text)
    except Exception as e:
        return f"Error reading DOCX: {e}"

def extract_text_from_xlsx(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            row_text = [str(cell) if cell is not None else "" for cell in row]
            rows.append("\t".join(row_text))
        return "\n".join(rows)
    except Exception as e:
        return f"Error reading XLSX: {e}"

def extract_text_from_pdf(filepath):
    try:
        reader = PyPDF2.PdfReader(filepath)
        full_text = [page.extract_text() for page in reader.pages]
        return "\n".join(full_text)
    except Exception as e:
        return f"Error reading PDF: {e}"

# ============================================================
# Helper function to flatten retrieved document items
# ============================================================
def flatten_item(item):
    if isinstance(item, str):
        return item
    elif isinstance(item, list):
        return " ".join(flatten_item(x) for x in item)
    else:
        return str(item)

# ============================================================
# Presentation Layer: RAG Tab (Previously Vectors and LLM)
# ============================================================
class VectorsLLMFrame(ttk.Frame):
    def __init__(self, container, *args, **kwargs):
        super().__init__(container, *args, **kwargs)
        # Initialize Chroma DB client (without persist_directory param)
        self.chroma_client = chromadb.Client()
        self.collection = self.chroma_client.get_or_create_collection(name="orinda_collection")
        
        # Title label (tab now named RAG)
        title_label = ttk.Label(self, text="RAG", font=BASE_FONT)
        title_label.pack(padx=10, pady=5, anchor="w")
        
        # Upload Files Section
        upload_frame = ttk.Frame(self)
        upload_frame.pack(fill="x", padx=10, pady=5)
        upload_button = ttk.Button(upload_frame, text="Upload Files", command=self.upload_files)
        upload_button.pack(side="left")
        self.upload_status = ttk.Label(upload_frame, text="", font=BASE_FONT)
        self.upload_status.pack(side="left", padx=10)
        
        # Listbox to show uploaded file names
        listbox_frame = ttk.Frame(self)
        listbox_frame.pack(fill="both", padx=10, pady=(0,10))
        listbox_label = ttk.Label(listbox_frame, text="Uploaded Files:", font=BASE_FONT)
        listbox_label.pack(anchor="w")
        self.uploaded_files_list = tk.Listbox(listbox_frame, font=BASE_FONT, height=5)
        self.uploaded_files_list.pack(fill="both", expand=True)
        
        # Query Section for RAG
        query_frame = ttk.Frame(self)
        query_frame.pack(fill="x", padx=10, pady=5)
        query_label = ttk.Label(query_frame, text="Enter Query for RAG:", font=BASE_FONT)
        query_label.pack(side="left")
        self.query_entry = ttk.Entry(query_frame, font=BASE_FONT)
        self.query_entry.pack(side="left", fill="x", expand=True, padx=5)
        rag_button = ttk.Button(query_frame, text="Perform RAG", command=self.perform_rag)
        rag_button.pack(side="left", padx=5)
        
        # Results Display Area
        self.results_display = ScrolledText(self, wrap=tk.WORD, height=15, font=BASE_FONT)
        self.results_display.pack(fill="both", expand=True, padx=10, pady=10)
    
    def upload_files(self):
        filetypes = [
            ("Word Files", "*.docx"),
            ("Excel Files", "*.xlsx"),
            ("PDF Files", "*.pdf")
        ]
        files = filedialog.askopenfilenames(title="Select Files", filetypes=filetypes)
        if not files:
            return
        count = 0
        for filepath in files:
            ext = os.path.splitext(filepath)[1].lower()
            text = ""
            if ext == ".docx":
                text = extract_text_from_docx(filepath)
            elif ext == ".xlsx":
                text = extract_text_from_xlsx(filepath)
            elif ext == ".pdf":
                text = extract_text_from_pdf(filepath)
            else:
                continue  # Skip unsupported types
            # Generate an embedding (dummy) from the extracted text.
            embedding = generate_embedding(text)
            # Use the filename as the document ID.
            file_id = os.path.basename(filepath)
            self.collection.add(documents=[text], ids=[file_id], embeddings=[embedding])
            count += 1
            # Add file name to listbox.
            self.uploaded_files_list.insert(tk.END, file_id)
        self.upload_status.config(text=f"Uploaded {count} file(s).")
    
    def perform_rag(self):
        """
        Perform Retrieval-Augmented Generation (RAG):
         1. Compute a query embedding.
         2. Query the Chroma DB collection for top related documents.
         3. Combine the retrieved context with the user's query and call the LLM.
        """
        query_text = self.query_entry.get().strip()
        if not query_text:
            messagebox.showerror("Input Error", "Please enter a query.")
            return
        # Compute query embedding (using dummy generator)
        query_embedding = generate_embedding(query_text)
        results = self.collection.query(query_embeddings=[query_embedding], n_results=3)
        context_parts = results.get("documents", [])
        flattened = [flatten_item(item) for item in context_parts]
        context = "\n\n".join(flattened)
        prompt = (f"Context:\n{context}\n\n"
                  f"User Query: {query_text}\n\n"
                  "Please provide a helpful answer based on the context.")
        rag_response = llm_chat(prompt)
        self.results_display.delete(1.0, tk.END)
        self.results_display.insert(tk.END, f"RAG Response:\n{rag_response}")

# ============================================================
# Presentation Layer: LLM Chat Tab
# ============================================================
class ChatFrame(ttk.Frame):
    def __init__(self, container, *args, **kwargs):
        super().__init__(container, *args, **kwargs)
        self.chat_display = ScrolledText(self, wrap=tk.WORD, height=20, font=BASE_FONT)
        self.chat_display.pack(fill="both", padx=10, pady=10, expand=True)
        input_frame = ttk.Frame(self)
        input_frame.pack(fill="x", padx=10, pady=5)
        self.entry = ttk.Entry(input_frame, font=BASE_FONT)
        self.entry.pack(side="left", fill="x", expand=True, padx=(0, 5))
        self.send_button = ttk.Button(input_frame, text="Send", command=self.send_message)
        self.send_button.pack(side="left", padx=(0, 5))
        manage_frame = ttk.Frame(self)
        manage_frame.pack(fill="x", padx=10, pady=5)
        self.save_button = ttk.Button(manage_frame, text="Save Chat", command=self.save_chat)
        self.save_button.pack(side="left", padx=(0, 5))
        self.load_button = ttk.Button(manage_frame, text="Load Chat", command=self.load_chat)
        self.load_button.pack(side="left", padx=(0, 5))
    
    def send_message(self):
        query = self.entry.get().strip()
        if not query:
            messagebox.showerror("Input Error", "Please enter a message before sending.")
            return
        self.chat_display.insert(tk.END, f"You: {query}\n")
        response = llm_chat(query)
        self.chat_display.insert(tk.END, f"LLM: {response}\n\n")
        self.entry.delete(0, tk.END)
    
    def save_chat(self):
        chat_text = self.chat_display.get(1.0, tk.END).strip()
        if not chat_text:
            messagebox.showinfo("Save Chat", "No chat history to save.")
            return
        file_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text Files", "*.txt"), ("All Files", "*.*")],
            title="Save Chat History"
        )
        if file_path:
            try:
                with open(file_path, "w", encoding="utf-8") as f:
                    f.write(chat_text)
                messagebox.showinfo("Save Chat", "Chat history saved successfully.")
            except Exception as e:
                messagebox.showerror("Save Chat Error", f"Failed to save chat:\n{e}")
    
    def load_chat(self):
        file_path = filedialog.askopenfilename(
            filetypes=[("Text Files", "*.txt"), ("All Files", "*.*")],
            title="Load Chat History"
        )
        if file_path:
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    chat_text = f.read()
                self.chat_display.delete(1.0, tk.END)
                self.chat_display.insert(tk.END, chat_text)
                messagebox.showinfo("Load Chat", "Chat history loaded successfully.")
            except Exception as e:
                messagebox.showerror("Load Chat Error", f"Failed to load chat:\n{e}")

# ============================================================
# Main Application: Orinda
# ============================================================
class OrindaApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Orinda")
        self.geometry("900x700")
        
        # Apply Park Theme Light from the themes folder.
        self.style = ttk.Style(self)
        try:
            theme_path = os.path.join(os.path.dirname(__file__), "themes", "park_light.tcl")
            self.tk.call("source", theme_path)
            self.style.theme_use("park_light")
        except tk.TclError as e:
            print("Could not load 'park_light' theme. Using default theme.", e)
        
        self.option_add("*Font", BASE_FONT)
        
        # Logo at the Top (sized 200x100).
        self.logo_frame = ttk.Frame(self)
        self.logo_frame.pack(padx=10, pady=10)
        logo_path = os.path.join(os.path.dirname(__file__), "ORINDA.png")
        if os.path.exists(logo_path):
            img = Image.open(logo_path)
            img = img.resize((200, 100), Image.Resampling.LANCZOS)
            self.logo_image = ImageTk.PhotoImage(img)
            self.logo_label = ttk.Label(self.logo_frame, image=self.logo_image)
            self.logo_label.pack()
        else:
            self.logo_label = ttk.Label(self.logo_frame, text="(Logo Not Found)", font=BASE_FONT)
            self.logo_label.pack()
        
        # Notebook Tabs
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)
        self.chat_tab = ChatFrame(self.notebook)
        self.notebook.add(self.chat_tab, text="LLM Chat")
        self.vectors_tab = VectorsLLMFrame(self.notebook)
        self.notebook.add(self.vectors_tab, text="RAG")
        
        # Exit Button
        exit_frame = ttk.Frame(self)
        exit_frame.pack(fill="x", padx=10, pady=10)
        exit_button = ttk.Button(exit_frame, text="Exit", command=self.exit_app)
        exit_button.pack(side="right")
    
    def exit_app(self):
        # Before exiting, clear the embeddings in Chroma.
        try:
            all_ids = self.vectors_tab.collection.get()['ids']
            if all_ids:
                self.vectors_tab.collection.delete(ids=all_ids)
                print("Chroma collection cleared.")
        except Exception as e:
            print(f"Error clearing Chroma collection: {e}")
        if messagebox.askokcancel("Exit", "Are you sure you want to exit?"):
            self.destroy()

if __name__ == "__main__":
    app = OrindaApp()
    app.mainloop()
